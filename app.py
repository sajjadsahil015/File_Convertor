import streamlit as st
import pandas as pd
import os
from io import BytesIO

# Fix for openpyxl PatternFill 'extLst' error with third-party Excel files
try:
    import openpyxl.styles.fills
    _orig_pattern_fill = openpyxl.styles.fills.PatternFill.__init__
    def _patched_pattern_fill(self, *args, **kwargs):
        kwargs.pop('extLst', None)
        return _orig_pattern_fill(self, *args, **kwargs)
    openpyxl.styles.fills.PatternFill.__init__ = _patched_pattern_fill
except Exception:
    pass

st.set_page_config(page_title="Data sweeper", layout='wide')
st.title("Data sweeper")
st.write("Transform your file between CSV and Excel formats with built-in data cleaning and visualization!")

uploaded_files = st.file_uploader("Upload your file(CSV or Excel):", type=["csv","xlsx"],accept_multiple_files=True)

if uploaded_files:
    for file in uploaded_files:
        file_ext = os.path.splitext(file.name)[-1].lower()

        try:
            if file_ext == ".csv":
                df = pd.read_csv(file)
            elif file_ext == ".xlsx":
                df = pd.read_excel(file)
            else:
                st.error(f"Unsupported file type: {file_ext}")
                continue
        except Exception as e:
            st.error(f"Error reading file '{file.name}': {e}")
            continue
        st.write(f"**File Name:** {file.name}")
        st.write(f"**File Size:** {file.size / 1024:.2f} KB")

        st.write("Preview the Head of the Dataframe")
        st.dataframe(df.head())

        st.subheader("Data cleaning options")
        if st.checkbox(f"Clean Data for {file.name}"):
            col1, col2 = st.columns(2)
            with col1:
                if st.button(f"Remove Duplicates from {file.name}"):
                    df.drop_duplicates(inplace=True)
                    st.success("Duplicates Removed!")
            with col2:
                if st.button(f"Fill missing values for {file.name}"):
                    numeric_col = df.select_dtypes(include=["number"]).columns
                    if not numeric_col.empty:
                        df[numeric_col] = df[numeric_col].fillna(df[numeric_col].mean())
                        st.success("Missing values have been filled!")
                    else:
                        st.info("No numeric columns found to fill missing values.")

        st.subheader("Select Columns to Convert")
        columns = st.multiselect(f"Choose Columns for {file.name}", df.columns, default=df.columns)
        if columns:
            df = df[columns]
        else:
            st.warning("Please select at least one column.")

        st.subheader("Data Visualization")
        if st.checkbox(f"Show visualization for {file.name}"):
            numeric_df = df.select_dtypes(include="number")
            if not numeric_df.empty:
                st.bar_chart(numeric_df.iloc[:, :2])
            else:
                st.info("No numeric columns found for visualization.")
        
        st.subheader("Conversion Options")
        conversion_type = st.radio(f"Convert {file.name} to:", ["CSV", "Excel"], key=f"radio_{file.name}")
        if st.button(f"Convert {file.name}", key=f"btn_{file.name}"):
            buffer = BytesIO()
            base_name = os.path.splitext(file.name)[0]
            if conversion_type == "CSV":
                df.to_csv(buffer, index=False)
                file_name = f"{base_name}.csv"
                mime_type = "text/csv"
            
            elif conversion_type == "Excel":
                df.to_excel(buffer, index=False, engine="openpyxl")
                file_name = f"{base_name}.xlsx"
                mime_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            buffer.seek(0)

            st.download_button(
                label=f"Download {file.name} as {conversion_type}",
                data=buffer,
                file_name=file_name,
                mime=mime_type,
                key=f"dl_{file.name}"
            )

            st.success("Processed successfully!")
        
        st.divider()


